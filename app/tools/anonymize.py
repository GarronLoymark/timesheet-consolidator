#!/usr/bin/env python3
"""Anonimiza los fixtures de timesheet: reemplaza nombres (pestañas y columna
Resource) por etiquetas genéricas consistentes con el roster, y las URLs
internas por un marcador. Conserva horas, fechas, Job Codes y estructura.

Salidas:
  app/test/fixtures/team.anon.xlsx
  app/test/fixtures/client.anon.xlsx
  app/config/config-inicial.json  (roster con nombres genéricos)
Imprime el mapeo real -> genérico para actualizar las pruebas.
"""
import json
import re
import unicodedata
from pathlib import Path

import openpyxl

APP = Path(__file__).resolve().parents[1]
FIX = APP / "test" / "fixtures"
CFG_PATH = APP / "config" / "config-inicial.json"


def key(s):
    s = unicodedata.normalize("NFD", str(s if s is not None else ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s).strip().lower()


cfg = json.loads(CFG_PATH.read_text(encoding="utf-8"))
roster = cfg["roster"]

# ---- Mapeo real -> genérico ----
display = {}  # realName -> generic
by_key = {}   # key(realName) -> generic
counters = {}


def gen_for(team, tipo):
    n = counters.get((team, tipo), 0) + 1
    counters[(team, tipo)] = n
    if tipo == "PM":
        return f"{team} PM"
    if tipo == "On demand":
        return f"{team} OnDem {n}"
    if tipo == "No incluir":
        return f"{team} {n}"
    return f"{team} Dev {n}"


for p in roster:
    g = gen_for(p.get("equipo") or "X", p.get("tipo") or "Fijo")
    display[p["dev"]] = g
    by_key[key(p["dev"])] = g

# PMs que solo aparecen en la columna pm (no son filas del roster)
for p in roster:
    pm = (p.get("pm") or "").strip()
    if pm and key(pm) not in by_key:
        g = f"{p.get('equipo') or 'X'} PM"
        display[pm] = g
        by_key[key(pm)] = g

externs = {}


def resolve_person(title):
    """Mapea el nombre de una hoja (abreviado) a su genérico, incluso vacía."""
    k = key(title).rstrip(".").strip()
    if k in by_key:
        return by_key[k]
    fn = k.split(" ")[0].rstrip(".")
    for real, g in display.items():
        if key(real).split(" ")[0] == fn:
            return g
    return None


def generic(name):
    k = key(name)
    if k in by_key:
        return by_key[k]
    if k == "":
        return name
    if k not in externs:
        externs[k] = f"Externo {len(externs) + 1}"
        by_key[k] = externs[k]
        display[name.strip()] = externs[k]
    return externs[k]


URL_RE = re.compile(r"https?://\S+")
url_map = {}


def repl_url(m):
    u = m.group(0)
    if u not in url_map:
        url_map[u] = f"https://example.invalid/anon/{len(url_map) + 1}"
    return url_map[u]


# Reemplazo de nombres completos en texto libre (de más largo a más corto)
names_sorted = sorted(display.keys(), key=lambda s: -len(s.strip()))


def scrub_text(val):
    if not isinstance(val, str):
        return val
    out = URL_RE.sub(repl_url, val)
    for real in names_sorted:
        r = real.strip()
        if r and r in out:
            out = out.replace(r, display[real])
    return out


def anon_workbook(src, dst, is_client=False):
    wb = openpyxl.load_workbook(src, data_only=True)
    renames = {}
    for ws in wb.worksheets:
        # localizar encabezado (client + hours) en las primeras 5 filas
        header_row = None
        res_col = None
        for i in range(1, min(5, ws.max_row) + 1):
            vals = {key(c.value): c.column for c in ws[i] if c.value is not None}
            if "client" in vals and "hours" in vals:
                header_row = i
                res_col = vals.get("resource")
                break
        # determinar persona dominante de la hoja por la columna Resource
        dominant = None
        if header_row and res_col:
            from collections import Counter

            cnt = Counter()
            for row in ws.iter_rows(min_row=header_row + 1, min_col=res_col, max_col=res_col, values_only=True):
                v = row[0]
                if isinstance(v, str) and v.strip():
                    cnt[v.strip()] += 1
            if cnt:
                dominant = cnt.most_common(1)[0][0]
        # reemplazar celdas de texto
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    if res_col and header_row and c.row > header_row and c.column == res_col and c.value.strip():
                        c.value = generic(c.value)
                    else:
                        c.value = scrub_text(c.value)
        # renombrar la hoja
        if dominant:
            renames[ws.title] = generic(dominant)
        else:
            r = resolve_person(ws.title)
            renames[ws.title] = r if r else ws.title
    # aplicar renombres evitando colisiones
    used = set()
    for ws in wb.worksheets:
        new = renames.get(ws.title, ws.title)
        base = new
        k = 2
        while new in used:
            new = f"{base} ({k})"
            k += 1
        used.add(new)
        ws.title = new[:31]
    wb.save(dst)


anon_workbook(FIX / "team.xlsx", FIX / "team.anon.xlsx")
anon_workbook(FIX / "client.xlsx", FIX / "client.anon.xlsx", is_client=True)

# ---- Config con roster genérico ----
for p in roster:
    p["dev"] = display.get(p["dev"], p["dev"])
    if p.get("pm"):
        p["pm"] = display.get(p["pm"], f"{p.get('equipo') or 'X'} PM")
CFG_PATH.write_text(json.dumps(cfg, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

print("=== MAPEO real -> genérico ===")
for real in sorted(display, key=lambda s: display[s]):
    print(f"{real!r:34} -> {display[real]}")
